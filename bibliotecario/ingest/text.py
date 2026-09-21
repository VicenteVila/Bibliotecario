"""Extracción de texto plano, Markdown/CSV y DOCX."""
from __future__ import annotations

from pathlib import Path

from bibliotecario.ingest.base import Document


def extract_txt(path: str | Path) -> Document:
    p = Path(path)
    text = p.read_text(encoding="utf-8", errors="ignore")
    if not text.strip():
        raise ValueError("Archivo de texto vacío")
    return Document(source=str(p), source_type="txt", title=p.stem, text=text)


def extract_docx(path: str | Path) -> Document:
    from docx import Document as Docx

    p = Path(path)
    d = Docx(str(p))
    parts = [para.text for para in d.paragraphs if para.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    if not text.strip():
        raise ValueError("DOCX sin texto extraíble")
    title = (d.core_properties.title or "").strip() or p.stem
    return Document(source=str(p), source_type="docx", title=title[:200], text=text)
