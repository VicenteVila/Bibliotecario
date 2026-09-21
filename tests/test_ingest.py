import numpy as np

from bibliotecario.ingest import base, citations, pipeline, quality, schemas
from bibliotecario.ingest.text import extract_docx, extract_txt
from bibliotecario.ingest.youtube import _parse_vtt, video_id


def test_chunk_text_solape():
    text = "\n".join(f"párrafo {i} " + "x" * 200 for i in range(20))
    chunks = base.chunk_text(text, size=500, overlap=50)
    assert len(chunks) > 1
    assert all(len(c) <= 500 + 50 + 1 for c in chunks)


def test_chunk_text_corto():
    assert base.chunk_text("hola") == ["hola"]
    assert base.chunk_text("") == []


def test_fingerprint_dedup():
    assert base.fingerprint("a") != base.fingerprint("b")


def test_extract_txt(tmp_path):
    f = tmp_path / "nota.txt"
    f.write_text("contenido de prueba " * 20)
    doc = extract_txt(f)
    assert doc.source_type == "txt" and len(doc.text) > 100


def test_extract_docx(tmp_path):
    from docx import Document as Docx
    f = tmp_path / "doc.docx"
    d = Docx()
    d.add_paragraph("párrafo importante " * 20)
    d.save(str(f))
    doc = extract_docx(f)
    assert doc.source_type == "docx" and "importante" in doc.text


def test_extract_pdf_generado(tmp_path):
    import pymupdf
    f = tmp_path / "t.pdf"
    d = pymupdf.open()
    page = d.new_page()
    page.insert_textbox(pymupdf.Rect(72, 72, 540, 750), "Texto digital del paper. " * 60)
    d.save(str(f))
    d.close()
    from bibliotecario.ingest.pdf import extract_pdf
    doc = extract_pdf(f)
    assert doc.source_type == "pdf" and "digital" in doc.text.lower()


def test_video_id():
    assert video_id("https://www.youtube.com/watch?v=abc123XYZ_-") == "abc123XYZ_-"
    assert video_id("https://youtu.be/abc123XYZ_-") == "abc123XYZ_-"
    assert video_id("https://example.com/x") is None


def test_parse_vtt():
    vtt = "WEBVTT\n\n00:00:01.000 --> 00:00:03.000\nhola mundo\n\n00:01:10.000 --> 00:01:12.000\nsegundo\n"
    out = _parse_vtt(vtt)
    assert out[0].startswith("[00:01]") or "[00:00]" in out[0]
    assert "hola mundo" in out[0]


def test_validators():
    good = base.Document(source="u", source_type="url", title="T", text="x" * 1000)
    assert schemas.validate(good)["valid"]
    bad = base.Document(source="u", source_type="url", title="", text="corto")
    assert not schemas.validate(bad)["valid"]


def test_quality():
    s, alerts = quality.score("url", "x" * 100)
    assert s < 100 and "contenido_muy_corto" in alerts
    s2, _ = quality.score("url", "x" * 5000)
    assert s2 == 100


def test_citations():
    c = citations.extract_citations("Ver arXiv:2608.06714 y https://example.com/a más doi 10.1234/abc.")
    assert "2608.06714" in c["arxiv"]
    assert c["urls"] == ["https://example.com/a"]


def test_store_dedup(tmp_path, monkeypatch):
    from bibliotecario.core import storage
    monkeypatch.setattr(storage, "db_path", lambda: tmp_path / "t.db")
    storage.init_db(tmp_path / "t.db")
    monkeypatch.setattr(pipeline, "encode", lambda chunks: np.zeros((len(chunks), 4), dtype="float32"))
    doc = base.Document(source="f", source_type="txt", title="Doc", text="contenido " * 200)
    r1 = pipeline.store(doc)
    assert r1.ok and r1.doc_id is not None and r1.chunks > 0
    r2 = pipeline.store(doc)
    assert r2.ok and r2.doc_id == r1.doc_id and "duplicado" in r2.quality_alerts[0]
